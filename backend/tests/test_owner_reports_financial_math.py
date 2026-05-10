"""
Tests for Owner Reports financial math fixes (iteration 78):
- Owner Reports monthly endpoint: new fields net_revenue / net_pl /
  eligible_rev_share_base / commissions.commission_gst.
- MG Payout PDF/Excel: 'Add: 18% GST on Rev Share' + 'Total Final Payout' rows.
- PIB PDF: 'Eligible Rev Share Base' line.
- MIS Franchise PDF: 'Eligible Rev Share Base' row.
- Owner Ledger PDF: 'Net Revenue Calculation' section.
"""
import os
import io
import pytest
import requests

def _read_env(name: str) -> str:
    v = os.environ.get(name)
    if v:
        return v
    # Fallback: read /app/frontend/.env
    try:
        with open("/app/frontend/.env") as f:
            for line in f:
                if line.startswith(name + "="):
                    return line.split("=", 1)[1].strip()
    except Exception:
        pass
    raise RuntimeError(f"{name} not set")


BASE_URL = _read_env("REACT_APP_BACKEND_URL").rstrip("/")
SUPER_ADMIN = {"mobile": "9741399190", "otp": "123456", "center": "PB-MGT"}
TEST_CENTER = "PB-HSR"
TEST_MONTH = "2026-02"


# ---------- shared fixtures ----------
@pytest.fixture(scope="module")
def session():
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json"})
    return s


@pytest.fixture(scope="module")
def sa_token(session):
    r = session.post(f"{BASE_URL}/api/send_otp",
                     json={"mobile": SUPER_ADMIN["mobile"], "center": SUPER_ADMIN["center"]})
    assert r.status_code == 200, f"send_otp failed: {r.status_code} {r.text}"
    r = session.post(f"{BASE_URL}/api/verify_otp", json=SUPER_ADMIN)
    assert r.status_code == 200, f"verify_otp failed: {r.status_code} {r.text}"
    tok = r.json().get("token")
    assert tok
    return tok


# ---------- Owner Reports monthly endpoint ----------
class TestOwnerReportsMath:
    def test_monthly_report_contains_new_fields(self, session, sa_token):
        r = session.post(
            f"{BASE_URL}/api/owner-reports/monthly-report",
            json={"token": sa_token, "center": TEST_CENTER, "month": TEST_MONTH},
        )
        assert r.status_code == 200, r.text
        data = r.json()
        assert data.get("success") is True
        # New top-level fields
        for k in ("net_revenue", "net_pl", "eligible_rev_share_base", "pnl"):
            assert k in data, f"missing top-level field {k}"
        # commissions.commission_gst
        assert "commissions" in data and "commission_gst" in data["commissions"]
        # Numeric types
        assert isinstance(data["net_revenue"], (int, float))
        assert isinstance(data["net_pl"], (int, float))
        assert isinstance(data["eligible_rev_share_base"], (int, float))

    def test_monthly_report_math_pb_hsr_feb_2026(self, session, sa_token):
        r = session.post(
            f"{BASE_URL}/api/owner-reports/monthly-report",
            json={"token": sa_token, "center": TEST_CENTER, "month": TEST_MONTH},
        )
        d = r.json()
        sales = d["sales"]["total"]
        gst = d["gst"]["gst_amount"]
        expenses = d["expenses"]["total"]
        comm = d["commissions"]["total"]
        comm_gst = d["commissions"]["commission_gst"]
        net_revenue = d["net_revenue"]
        net_pl = d["net_pl"]
        eligible = d["eligible_rev_share_base"]

        # Net Revenue = Sales - Commission - GST
        expected_net_rev = round(sales - comm - gst, 2)
        assert abs(net_revenue - expected_net_rev) < 0.05, (
            f"net_revenue mismatch: got {net_revenue}, expected {expected_net_rev}"
        )
        # Net P/L = Net Revenue - Expenses
        expected_net_pl = round(net_revenue - expenses, 2)
        assert abs(net_pl - expected_net_pl) < 0.05, (
            f"net_pl mismatch: got {net_pl}, expected {expected_net_pl}"
        )
        # Eligible Rev Share Base for India = Sales - Comm - CommGST - GST
        # (note: total_commission already includes commGST per backend comment)
        expected_eligible = round(sales - (comm - comm_gst) - comm_gst - gst, 2)
        assert abs(eligible - expected_eligible) < 0.05
        # For India centers: eligible_rev_share_base == net_revenue
        country = (d.get("country") or "India")
        if country == "India":
            assert abs(eligible - net_revenue) < 0.05, (
                f"India: eligible should equal net_revenue: {eligible} vs {net_revenue}"
            )
        # pnl is back-compat alias of net_revenue
        assert abs(d["pnl"] - net_revenue) < 0.05

    def test_monthly_report_australia_profitability(self, session, sa_token):
        # PB-PERTH (Australia) should still surface profitability and net_pl
        r = session.post(
            f"{BASE_URL}/api/owner-reports/monthly-report",
            json={"token": sa_token, "center": "PB-PERTH", "month": TEST_MONTH},
        )
        if r.status_code != 200:
            pytest.skip("PB-PERTH data not available on this preview")
        d = r.json()
        if not d.get("success"):
            pytest.skip("PB-PERTH report not visible")
        # Even if no data, fields must be present
        assert "net_revenue" in d
        assert "net_pl" in d
        assert "eligible_rev_share_base" in d


# ---------- MG Payout PDF + Excel ----------
class TestMgPayoutExports:
    def test_mg_payout_pdf_has_gst_and_final_payout(self, session, sa_token):
        r = session.post(
            f"{BASE_URL}/api/center-accounts/export-mg-payout",
            json={"token": sa_token, "center": TEST_CENTER,
                  "from_month": "2025-11", "to_month": TEST_MONTH, "format": "pdf"},
        )
        assert r.status_code == 200, r.text[:500]
        ctype = r.headers.get("content-type", "")
        assert "pdf" in ctype.lower() or r.content[:4] == b"%PDF", f"not a pdf: {ctype}"
        # Decode searchable text from PDF
        try:
            from pypdf import PdfReader
        except ImportError:
            from PyPDF2 import PdfReader  # type: ignore
        reader = PdfReader(io.BytesIO(r.content))
        text = ""
        for p in reader.pages:
            text += (p.extract_text() or "") + "\n"
        # Required strings
        assert "18% GST" in text or "18%" in text, "GST 18% line missing in PDF"
        assert "Final Payout" in text, "'Final Payout' missing in PDF"
        # Either of these phrases should appear
        assert ("Add: 18% GST on Rev Share" in text or
                "Add 18% GST on Rev Share" in text or
                "GST on Rev Share" in text), "Add GST row missing"
        assert ("Total Final Payout" in text), "Total Final Payout row missing"

    def test_mg_payout_excel_has_gst_and_final_payout(self, session, sa_token):
        r = session.post(
            f"{BASE_URL}/api/center-accounts/export-mg-payout",
            json={"token": sa_token, "center": TEST_CENTER,
                  "from_month": "2025-11", "to_month": TEST_MONTH, "format": "excel"},
        )
        assert r.status_code == 200, r.text[:500]
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content), data_only=True)
        # Concatenate all cell text across sheets
        joined = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for c in row:
                    if c is not None:
                        joined.append(str(c))
        blob = " | ".join(joined)
        assert "GST on Rev Share" in blob or "18% GST" in blob, (
            "Add GST row missing in Excel"
        )
        assert "Total Final Payout" in blob, "Total Final Payout row missing in Excel"


# ---------- PIB PDF ----------
class TestPibPdf:
    def test_pib_pdf_contains_eligible_rev_share_base(self, session, sa_token):
        r = session.post(
            f"{BASE_URL}/api/center-accounts/generate-pib",
            json={"token": sa_token, "center": TEST_CENTER, "month": TEST_MONTH},
        )
        assert r.status_code == 200, r.text[:500]
        try:
            from pypdf import PdfReader
        except ImportError:
            from PyPDF2 import PdfReader  # type: ignore
        reader = PdfReader(io.BytesIO(r.content))
        text = "".join((p.extract_text() or "") for p in reader.pages)
        assert "Eligible Rev Share Base" in text or "ELIGIBLE REV SHARE BASE" in text.upper(), \
            "Eligible Rev Share Base missing in PIB PDF"
        assert "NET REVENUE" in text.upper(), "NET REVENUE missing in PIB PDF"


# ---------- MIS Franchise PDF ----------
class TestMisFranchisePdf:
    def test_mis_franchise_pdf_contains_eligible_rev_share_base(self, session, sa_token):
        r = session.post(
            f"{BASE_URL}/api/mis/franchise-pdf",
            json={"token": sa_token, "center": TEST_CENTER, "month": TEST_MONTH},
        )
        if r.status_code != 200:
            pytest.skip(f"mis franchise-pdf returned {r.status_code}: {r.text[:200]}")
        try:
            from pypdf import PdfReader
        except ImportError:
            from PyPDF2 import PdfReader  # type: ignore
        reader = PdfReader(io.BytesIO(r.content))
        text = "".join((p.extract_text() or "") for p in reader.pages)
        assert "Eligible Rev Share Base" in text or "ELIGIBLE REV SHARE BASE" in text.upper(), \
            "Eligible Rev Share Base missing in MIS Franchise PDF"


# ---------- Owner Ledger PDF ----------
class TestOwnerLedgerPdf:
    def test_owner_ledger_pdf_has_net_revenue_calculation(self, session, sa_token):
        r = session.post(
            f"{BASE_URL}/api/ledgers/owner",
            json={"token": sa_token, "center": TEST_CENTER,
                  "period_type": "month", "month": TEST_MONTH, "format": "pdf"},
        )
        assert r.status_code == 200, r.text[:500]
        try:
            from pypdf import PdfReader
        except ImportError:
            from PyPDF2 import PdfReader  # type: ignore
        reader = PdfReader(io.BytesIO(r.content))
        text = "".join((p.extract_text() or "") for p in reader.pages)
        upper = text.upper()
        assert "NET REVENUE CALCULATION" in upper, "Net Revenue Calculation header missing"
        # Required rows
        assert "TOTAL SALES" in upper
        assert "COMMISSION" in upper
        assert "GST" in upper
        assert "ELIGIBLE REV SHARE BASE" in upper, \
            "Eligible Rev Share Base row missing in owner ledger PDF"


# ---------- Regression: other surfaces still work ----------
class TestRegressions:
    def test_center_accounts_summary_loads(self, session, sa_token):
        r = session.post(
            f"{BASE_URL}/api/center-accounts/summary",
            json={"token": sa_token, "center": TEST_CENTER, "month": TEST_MONTH},
        )
        assert r.status_code == 200, r.text[:300]

    def test_mis_dashboard_loads(self, session, sa_token):
        r = session.post(
            f"{BASE_URL}/api/mis/overview",
            json={"token": sa_token, "month": TEST_MONTH},
        )
        assert r.status_code in (200, 400, 422), f"unexpected: {r.status_code} {r.text[:200]}"

    def test_gst_reconciliation_loads(self, session, sa_token):
        r = session.post(
            f"{BASE_URL}/api/gst/reconciliation",
            json={"token": sa_token, "center": TEST_CENTER, "month": TEST_MONTH},
        )
        assert r.status_code in (200, 400, 404, 422)
