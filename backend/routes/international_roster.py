# =======================================
# International Weekly Roster Module
# Master-driven, no hardcoding, scalable
# =======================================

from fastapi import APIRouter, HTTPException, Body
from pydantic import BaseModel
from typing import Optional, List
from datetime import datetime, timezone, timedelta, date as _date
import logging
import uuid
import calendar
import io
import urllib.parse

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api/international-roster", tags=["International Roster"])

db = None
verify_token = None

def set_db(database):
    global db
    db = database

def set_verify_token(func):
    global verify_token
    verify_token = func


# =======================================
# ACCESS HELPERS
# =======================================

async def check_access(token: str):
    session = verify_token(token)
    if not session:
        raise HTTPException(401, "Invalid or expired token")
    return session

def get_user_center(session):
    return (session.get("center") or "").upper()

def is_admin(session):
    return session.get("is_super_admin") or session.get("is_admin")

async def check_center_access(session, center_code):
    """Verify user has access to this international center."""
    if is_admin(session):
        return True
    user_center = get_user_center(session)
    if user_center == center_code.upper():
        return True
    raise HTTPException(403, "Access denied to this center")


# =======================================
# WEEK HELPERS
# =======================================

def get_week_info(target_date=None):
    """Get week info for a date. Week = Mon-Sun."""
    if target_date is None:
        target_date = _date.today()
    elif isinstance(target_date, str):
        target_date = _date.fromisoformat(target_date)
    
    # Find Monday of this week
    monday = target_date - timedelta(days=target_date.weekday())
    sunday = monday + timedelta(days=6)
    
    # ISO week number
    week_number = monday.isocalendar()[1]
    year = monday.year
    
    days = []
    for i in range(7):
        d = monday + timedelta(days=i)
        days.append({
            "date": d.isoformat(),
            "day": calendar.day_name[d.weekday()],
            "day_short": calendar.day_abbr[d.weekday()],
        })
    
    return {
        "week_number": week_number,
        "year": year,
        "week_start": monday.isoformat(),
        "week_end": sunday.isoformat(),
        "days": days,
        "label": f"Week {week_number} — {monday.strftime('%d %b')} to {sunday.strftime('%d %b %Y')}"
    }

def get_next_week_info():
    """Get next week's info (the one to plan on Sunday)."""
    next_monday = _date.today() + timedelta(days=(7 - _date.today().weekday()))
    return get_week_info(next_monday)


# =======================================
# MASTER ENDPOINTS
# =======================================

@router.post("/masters/list")
async def list_masters(req: dict = Body(...)):
    """List master items (roles, shifts, settings) for a center."""
    session = await check_access(req.get("token"))
    center = (req.get("center") or "").upper()
    master_type = req.get("type", "role")  # role, shift, settings
    
    query = {"type": master_type, "$or": [{"center": center}, {"center": ""}, {"center": "GLOBAL"}]}
    items = await db.roster_masters.find(query, {"_id": 0}).sort("order", 1).to_list(200)
    return {"success": True, "items": items, "type": master_type}


@router.post("/masters/save")
async def save_master(req: dict = Body(...)):
    """Create or update a master item."""
    session = await check_access(req.get("token"))
    if not is_admin(session):
        raise HTTPException(403, "Admin access required")
    
    item_id = req.get("item_id") or f"RM-{uuid.uuid4().hex[:8].upper()}"
    doc = {
        "item_id": item_id,
        "type": req.get("type", "role"),
        "center": (req.get("center") or "GLOBAL").upper(),
        "name": req.get("name", ""),
        "config": req.get("config", {}),
        "active": req.get("active", True),
        "order": req.get("order", 0),
        "updated_by": session.get("managerName", ""),
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }
    
    await db.roster_masters.update_one(
        {"item_id": item_id},
        {"$set": doc, "$setOnInsert": {"created_at": datetime.now(timezone.utc).isoformat()}},
        upsert=True
    )
    return {"success": True, "item_id": item_id, "message": "Master item saved"}


@router.post("/masters/delete")
async def delete_master(req: dict = Body(...)):
    session = await check_access(req.get("token"))
    if not is_admin(session):
        raise HTTPException(403, "Admin access required")
    await db.roster_masters.delete_one({"item_id": req.get("item_id")})
    return {"success": True, "message": "Deleted"}


@router.post("/masters/seed")
async def seed_masters(req: dict = Body(...)):
    """Seed default roles and shifts for a center."""
    session = await check_access(req.get("token"))
    if not is_admin(session):
        raise HTTPException(403, "Admin access required")
    
    center = (req.get("center") or "GLOBAL").upper()
    now = datetime.now(timezone.utc).isoformat()
    
    default_roles = ["Biller", "Plater", "Service", "Kitchen Hand", "Housekeeper", "Cleaner", "Runner", "Supervisor", "Chef", "Manager"]
    default_shifts = [
        {"name": "Morning", "config": {"default_in": "06:00", "default_out": "14:00"}},
        {"name": "Evening", "config": {"default_in": "14:00", "default_out": "22:00"}},
        {"name": "Full Day", "config": {"default_in": "09:00", "default_out": "21:00"}},
        {"name": "Split Shift", "config": {"default_in": "10:00", "default_out": "20:00", "break_minutes": 120}},
    ]
    
    seeded = 0
    for i, role in enumerate(default_roles):
        exists = await db.roster_masters.find_one({"type": "role", "name": role, "center": {"$in": [center, "GLOBAL"]}})
        if not exists:
            await db.roster_masters.insert_one({
                "item_id": f"RM-{uuid.uuid4().hex[:8].upper()}", "type": "role",
                "center": center, "name": role, "config": {}, "active": True,
                "order": i, "created_at": now, "updated_at": now
            })
            seeded += 1
    
    for i, shift in enumerate(default_shifts):
        exists = await db.roster_masters.find_one({"type": "shift", "name": shift["name"], "center": {"$in": [center, "GLOBAL"]}})
        if not exists:
            await db.roster_masters.insert_one({
                "item_id": f"RM-{uuid.uuid4().hex[:8].upper()}", "type": "shift",
                "center": center, "name": shift["name"], "config": shift["config"],
                "active": True, "order": i, "created_at": now, "updated_at": now
            })
            seeded += 1
    
    # Seed default settings
    exists = await db.roster_masters.find_one({"type": "settings", "center": center})
    if not exists:
        await db.roster_masters.insert_one({
            "item_id": f"RM-{uuid.uuid4().hex[:8].upper()}", "type": "settings",
            "center": center, "name": "Roster Settings",
            "config": {
                "fill_day": "Sunday", "week_start": "Monday", "week_end": "Sunday",
                "max_hours_per_week": 38, "max_hours_per_day": 12,
                "whatsapp_template": "Hi {employee_name},\n\nYou have been rostered for duty:\n\nCenter: {center}\nDate: {date} ({day})\nShift: {shift_type}\nTime: {in_time} - {out_time}\nRole: {duty_role}\n\nPlease confirm by replying YES or DENY.\n\nThank you,\n{center} Management",
                "reminder_hours": [6, 12],
                "confirmation_cutoff_hours": 24,
            },
            "active": True, "order": 0, "created_at": now, "updated_at": now
        })
        seeded += 1
    
    return {"success": True, "message": f"Seeded {seeded} master items for {center}"}


# =======================================
# EMPLOYEES FOR ROSTER
# =======================================

@router.post("/employees")
async def get_roster_employees(req: dict = Body(...)):
    """Get active employees for an international center."""
    session = await check_access(req.get("token"))
    center = (req.get("center") or "").upper()
    await check_center_access(session, center)
    
    employees = await db.employees.find(
        {"center": center, "status": {"$ne": "inactive"}},
        {"_id": 0, "name": 1, "mobile": 1, "whatsapp": 1, "designation": 1,
         "employeeCode": 1, "category": 1, "hourlyRate": 1, "target_takehome_rate": 1,
         "center": 1, "status": 1}
    ).sort("name", 1).to_list(500)
    
    # Normalize — use mobile as whatsapp fallback
    for emp in employees:
        emp["whatsapp"] = emp.get("whatsapp") or emp.get("mobile") or ""
        emp["hourly_rate"] = emp.get("target_takehome_rate") or emp.get("hourlyRate") or 0
    
    return {"success": True, "employees": employees, "count": len(employees)}


# =======================================
# WEEK INFO
# =======================================

@router.post("/week-info")
async def get_week_info_endpoint(req: dict = Body(...)):
    """Get week info for current or specific date."""
    await check_access(req.get("token"))
    target = req.get("date")
    if target:
        info = get_week_info(target)
    else:
        info = get_week_info()
    
    next_week = get_next_week_info()
    return {"success": True, "current_week": info, "next_week": next_week}


# =======================================
# ROSTER CRUD
# =======================================

@router.post("/roster/get")
async def get_roster(req: dict = Body(...)):
    """Get a roster for a specific week and center."""
    session = await check_access(req.get("token"))
    center = (req.get("center") or "").upper()
    week_start = req.get("week_start")
    await check_center_access(session, center)
    
    if not week_start:
        raise HTTPException(400, "week_start is required")
    
    # Get or calculate week info
    week_info = get_week_info(week_start)
    
    # Find roster header
    roster = await db.roster_weeks.find_one(
        {"center": center, "week_start": week_info["week_start"]},
        {"_id": 0}
    )
    
    # Get lines
    lines = []
    if roster:
        lines = await db.roster_lines.find(
            {"roster_id": roster["roster_id"]},
            {"_id": 0}
        ).sort([("date", 1), ("planned_in", 1)]).to_list(500)
    
    return {
        "success": True,
        "roster": roster,
        "lines": lines,
        "week_info": week_info,
        "exists": roster is not None
    }


@router.post("/roster/save")
async def save_roster(req: dict = Body(...)):
    """Create or update a roster week + lines."""
    session = await check_access(req.get("token"))
    center = (req.get("center") or "").upper()
    await check_center_access(session, center)
    
    week_start = req.get("week_start")
    if not week_start:
        raise HTTPException(400, "week_start is required")
    
    week_info = get_week_info(week_start)
    roster_id = req.get("roster_id") or f"RST-{center}-{week_info['week_start']}"
    now = datetime.now(timezone.utc).isoformat()
    user = session.get("managerName", "Unknown")
    
    # Upsert roster header
    existing = await db.roster_weeks.find_one({"roster_id": roster_id}, {"_id": 0})
    
    header = {
        "roster_id": roster_id,
        "center": center,
        "week_number": week_info["week_number"],
        "year": week_info["year"],
        "week_start": week_info["week_start"],
        "week_end": week_info["week_end"],
        "week_label": week_info["label"],
        "status": req.get("status") or (existing.get("status") if existing else "draft"),
        "updated_by": user,
        "updated_at": now,
    }
    
    if not existing:
        header["created_by"] = user
        header["created_at"] = now
        await db.roster_weeks.insert_one(header)
    else:
        await db.roster_weeks.update_one(
            {"roster_id": roster_id},
            {"$set": header}
        )
    
    # Save lines
    lines = req.get("lines") or []
    saved_ids = []
    
    for line in lines:
        line_id = line.get("line_id") or f"RSL-{uuid.uuid4().hex[:10].upper()}"
        line_date = line.get("date", "")
        
        # Auto-calc working hours
        planned_in = line.get("planned_in", "")
        planned_out = line.get("planned_out", "")
        break_min = int(line.get("break_minutes") or 0)
        working_hours = 0
        if planned_in and planned_out:
            try:
                t_in = datetime.strptime(planned_in, "%H:%M")
                t_out = datetime.strptime(planned_out, "%H:%M")
                diff = (t_out - t_in).total_seconds() / 3600
                if diff < 0:
                    diff += 24
                working_hours = round(diff - break_min / 60, 2)
            except ValueError:
                pass
        
        # Auto-derive day name
        day_name = ""
        if line_date:
            try:
                d = _date.fromisoformat(line_date)
                day_name = calendar.day_name[d.weekday()]
            except ValueError:
                pass
        
        line_doc = {
            "line_id": line_id,
            "roster_id": roster_id,
            "center": center,
            "date": line_date,
            "day": day_name,
            "week_number": week_info["week_number"],
            "shift_type": line.get("shift_type", ""),
            "planned_in": planned_in,
            "planned_out": planned_out,
            "break_minutes": break_min,
            "working_hours": working_hours,
            "employee_name": line.get("employee_name", ""),
            "employee_id": line.get("employee_id", ""),
            "whatsapp": line.get("whatsapp", ""),
            "duty_role": line.get("duty_role", ""),
            "backup_employee": line.get("backup_employee", ""),
            "notes": line.get("notes", ""),
            "status": line.get("status") or "draft",
            "confirmation_sent_at": line.get("confirmation_sent_at") or "",
            "response_at": line.get("response_at") or "",
            "response_note": line.get("response_note") or "",
            "replaced_by": line.get("replaced_by") or "",
            "replacement_history": line.get("replacement_history") or [],
            "attendance_synced": line.get("attendance_synced", False),
            "attendance_sync_at": line.get("attendance_sync_at") or "",
            "updated_by": user,
            "updated_at": now,
        }
        
        existing_line = await db.roster_lines.find_one({"line_id": line_id})
        if existing_line:
            await db.roster_lines.update_one({"line_id": line_id}, {"$set": line_doc})
        else:
            line_doc["created_at"] = now
            line_doc["created_by"] = user
            await db.roster_lines.insert_one(line_doc)
        saved_ids.append(line_id)
    
    # Log audit
    await db.roster_audit.insert_one({
        "roster_id": roster_id, "action": "saved",
        "details": f"Saved {len(lines)} lines", "by": user, "at": now, "_id_skip": True
    })
    
    return {"success": True, "roster_id": roster_id, "saved_lines": len(saved_ids), "message": "Roster saved"}


@router.post("/roster/delete-line")
async def delete_roster_line(req: dict = Body(...)):
    """Delete a single roster line."""
    session = await check_access(req.get("token"))
    line_id = req.get("line_id")
    if not line_id:
        raise HTTPException(400, "line_id required")
    
    line = await db.roster_lines.find_one({"line_id": line_id}, {"_id": 0})
    if line and line.get("status") in ("confirmed", "locked"):
        raise HTTPException(403, "Cannot delete confirmed/locked lines")
    
    await db.roster_lines.delete_one({"line_id": line_id})
    
    user = session.get("managerName", "Unknown")
    await db.roster_audit.insert_one({
        "roster_id": line.get("roster_id") if line else "",
        "action": "line_deleted", "details": f"Deleted line {line_id}",
        "by": user, "at": datetime.now(timezone.utc).isoformat()
    })
    return {"success": True, "message": "Line deleted"}


# =======================================
# COPY LAST WEEK
# =======================================

@router.post("/roster/copy-week")
async def copy_last_week(req: dict = Body(...)):
    """Copy previous week's roster as draft for target week."""
    session = await check_access(req.get("token"))
    center = (req.get("center") or "").upper()
    target_week_start = req.get("week_start")
    await check_center_access(session, center)
    
    if not target_week_start:
        raise HTTPException(400, "week_start required")
    
    target_info = get_week_info(target_week_start)
    # Find previous week
    prev_monday = _date.fromisoformat(target_info["week_start"]) - timedelta(days=7)
    prev_info = get_week_info(prev_monday)
    
    prev_roster = await db.roster_weeks.find_one(
        {"center": center, "week_start": prev_info["week_start"]},
        {"_id": 0}
    )
    if not prev_roster:
        raise HTTPException(404, "No previous week roster found to copy")
    
    prev_lines = await db.roster_lines.find(
        {"roster_id": prev_roster["roster_id"]},
        {"_id": 0}
    ).to_list(500)
    
    if not prev_lines:
        raise HTTPException(404, "Previous week roster has no lines")
    
    # Map old dates to new dates (same day of week)
    date_map = {}
    for i, day in enumerate(prev_info["days"]):
        date_map[day["date"]] = target_info["days"][i]["date"]
    
    # Create new lines
    new_lines = []
    for line in prev_lines:
        new_date = date_map.get(line["date"], "")
        if not new_date:
            continue
        new_lines.append({
            "date": new_date,
            "shift_type": line.get("shift_type", ""),
            "planned_in": line.get("planned_in", ""),
            "planned_out": line.get("planned_out", ""),
            "break_minutes": line.get("break_minutes", 0),
            "employee_name": line.get("employee_name", ""),
            "employee_id": line.get("employee_id", ""),
            "whatsapp": line.get("whatsapp", ""),
            "duty_role": line.get("duty_role", ""),
            "notes": "",
            "status": "draft",
        })
    
    # Save via the save endpoint logic
    save_req = {
        "token": req.get("token"),
        "center": center,
        "week_start": target_info["week_start"],
        "lines": new_lines
    }
    result = await save_roster(save_req)
    result["copied_from"] = prev_info["label"]
    result["message"] = f"Copied {len(new_lines)} lines from {prev_info['label']}"
    return result


# =======================================
# WHATSAPP CONFIRMATION
# =======================================

@router.post("/roster/send-confirmation")
async def send_confirmation(req: dict = Body(...)):
    """Generate WhatsApp links for pending roster lines and mark as sent."""
    session = await check_access(req.get("token"))
    center = (req.get("center") or "").upper()
    roster_id = req.get("roster_id")
    line_ids = req.get("line_ids")  # Optional: send only specific lines
    await check_center_access(session, center)
    
    if not roster_id:
        raise HTTPException(400, "roster_id required")
    
    # Get settings for message template
    settings = await db.roster_masters.find_one(
        {"type": "settings", "center": {"$in": [center, "GLOBAL"]}},
        {"_id": 0}
    )
    template = (settings.get("config", {}).get("whatsapp_template", "") if settings else
                "Hi {employee_name}, you are rostered at {center} on {date} ({day}) from {in_time} to {out_time} as {duty_role}. Please confirm YES or DENY.")
    
    # Get lines to send
    query = {"roster_id": roster_id, "status": {"$in": ["draft", "sent"]}}
    if line_ids:
        query["line_id"] = {"$in": line_ids}
    
    lines = await db.roster_lines.find(query, {"_id": 0}).to_list(500)
    
    if not lines:
        raise HTTPException(404, "No draft/pending lines to send")
    
    now = datetime.now(timezone.utc).isoformat()
    user = session.get("managerName", "Unknown")
    whatsapp_links = []
    
    for line in lines:
        phone = (line.get("whatsapp") or "").strip()
        if not phone:
            continue
        
        # Clean phone number
        phone_clean = phone.replace(" ", "").replace("-", "").replace("(", "").replace(")", "")
        if not phone_clean.startswith("+"):
            phone_clean = f"+{phone_clean}"
        
        # Build message from template
        message = template.format(
            employee_name=line.get("employee_name", ""),
            center=center,
            date=line.get("date", ""),
            day=line.get("day", ""),
            shift_type=line.get("shift_type", ""),
            in_time=line.get("planned_in", ""),
            out_time=line.get("planned_out", ""),
            duty_role=line.get("duty_role", ""),
        )
        
        wa_link = f"https://wa.me/{phone_clean.lstrip('+')}?text={urllib.parse.quote(message)}"
        
        # Update line status
        await db.roster_lines.update_one(
            {"line_id": line["line_id"]},
            {"$set": {"status": "sent", "confirmation_sent_at": now, "updated_by": user, "updated_at": now}}
        )
        
        whatsapp_links.append({
            "line_id": line["line_id"],
            "employee_name": line.get("employee_name"),
            "whatsapp": phone,
            "wa_link": wa_link,
            "date": line.get("date"),
            "shift": line.get("shift_type"),
        })
    
    # Update roster header status
    await db.roster_weeks.update_one(
        {"roster_id": roster_id},
        {"$set": {"status": "sent", "updated_by": user, "updated_at": now}}
    )
    
    # Audit
    await db.roster_audit.insert_one({
        "roster_id": roster_id, "action": "confirmation_sent",
        "details": f"Sent {len(whatsapp_links)} WhatsApp confirmations",
        "by": user, "at": now
    })
    
    return {"success": True, "links": whatsapp_links, "count": len(whatsapp_links)}


@router.post("/roster/update-response")
async def update_response(req: dict = Body(...)):
    """Record employee's confirm/deny response."""
    session = await check_access(req.get("token"))
    line_id = req.get("line_id")
    response = req.get("response", "").lower()  # "confirmed" or "denied"
    note = req.get("note", "")
    
    if not line_id or response not in ("confirmed", "denied"):
        raise HTTPException(400, "line_id and response (confirmed/denied) required")
    
    now = datetime.now(timezone.utc).isoformat()
    user = session.get("managerName", "Unknown")
    
    line = await db.roster_lines.find_one({"line_id": line_id}, {"_id": 0})
    if not line:
        raise HTTPException(404, "Line not found")
    
    update = {
        "status": response,
        "response_at": now,
        "response_note": note,
        "updated_by": user,
        "updated_at": now,
    }
    
    await db.roster_lines.update_one({"line_id": line_id}, {"$set": update})
    
    # Update roster header status
    roster_id = line.get("roster_id")
    if roster_id:
        all_lines = await db.roster_lines.find({"roster_id": roster_id}, {"_id": 0, "status": 1}).to_list(500)
        statuses = set(l.get("status") for l in all_lines)
        if statuses == {"confirmed"}:
            new_status = "fully_confirmed"
        elif "confirmed" in statuses and ("sent" in statuses or "denied" in statuses or "draft" in statuses):
            new_status = "partially_confirmed"
        else:
            new_status = "sent"
        await db.roster_weeks.update_one({"roster_id": roster_id}, {"$set": {"status": new_status, "updated_at": now}})
    
    # Audit
    await db.roster_audit.insert_one({
        "roster_id": roster_id, "action": f"response_{response}",
        "details": f"{line.get('employee_name')} {response} for {line.get('date')} {line.get('shift_type')}. Note: {note}",
        "by": user, "at": now
    })
    
    return {"success": True, "message": f"Response recorded: {response}"}


@router.post("/roster/replace-employee")
async def replace_employee(req: dict = Body(...)):
    """Replace a denied employee with a new one."""
    session = await check_access(req.get("token"))
    line_id = req.get("line_id")
    new_employee = req.get("new_employee_name", "")
    new_whatsapp = req.get("new_whatsapp", "")
    
    if not line_id or not new_employee:
        raise HTTPException(400, "line_id and new_employee_name required")
    
    now = datetime.now(timezone.utc).isoformat()
    user = session.get("managerName", "Unknown")
    
    line = await db.roster_lines.find_one({"line_id": line_id}, {"_id": 0})
    if not line:
        raise HTTPException(404, "Line not found")
    
    # Build replacement history
    history = line.get("replacement_history") or []
    history.append({
        "original_employee": line.get("employee_name"),
        "original_whatsapp": line.get("whatsapp"),
        "denied_at": line.get("response_at"),
        "replaced_by": new_employee,
        "replaced_at": now,
        "replaced_by_user": user,
    })
    
    update = {
        "employee_name": new_employee,
        "whatsapp": new_whatsapp,
        "status": "draft",  # Reset to draft so new confirmation can be sent
        "replaced_by": new_employee,
        "replacement_history": history,
        "response_at": "",
        "response_note": "",
        "confirmation_sent_at": "",
        "updated_by": user,
        "updated_at": now,
    }
    
    await db.roster_lines.update_one({"line_id": line_id}, {"$set": update})
    
    await db.roster_audit.insert_one({
        "roster_id": line.get("roster_id"), "action": "replacement",
        "details": f"Replaced {line.get('employee_name')} with {new_employee} for {line.get('date')}",
        "by": user, "at": now
    })
    
    return {"success": True, "message": f"Replaced with {new_employee}. Send confirmation to proceed."}


# =======================================
# ATTENDANCE SYNC
# =======================================

@router.post("/roster/sync-attendance")
async def sync_attendance(req: dict = Body(...)):
    """Sync confirmed roster lines into international attendance."""
    session = await check_access(req.get("token"))
    roster_id = req.get("roster_id")
    if not roster_id:
        raise HTTPException(400, "roster_id required")
    
    now = datetime.now(timezone.utc).isoformat()
    user = session.get("managerName", "Unknown")
    
    lines = await db.roster_lines.find(
        {"roster_id": roster_id, "status": "confirmed", "attendance_synced": {"$ne": True}},
        {"_id": 0}
    ).to_list(500)
    
    if not lines:
        raise HTTPException(404, "No confirmed un-synced lines to process")
    
    synced = 0
    for line in lines:
        # Create/update international attendance record
        att_doc = {
            "employee_id": line.get("employee_name", "").upper(),
            "center": line.get("center"),
            "date": line.get("date"),
            "hours_worked": line.get("working_hours", 0),
            "week_number": line.get("week_number"),
            "month": int(line["date"].split("-")[1]) if line.get("date") else 0,
            "year": int(line["date"].split("-")[0]) if line.get("date") else 0,
            "source": "weekly_roster",
            "roster_line_id": line["line_id"],
            "duty_role": line.get("duty_role"),
            "shift_type": line.get("shift_type"),
            "planned_in": line.get("planned_in"),
            "planned_out": line.get("planned_out"),
            "updated_at": now,
            "updated_by": user,
        }
        
        await db.international_attendance.update_one(
            {"employee_id": att_doc["employee_id"], "center": att_doc["center"], "date": att_doc["date"]},
            {"$set": att_doc, "$setOnInsert": {"created_at": now, "created_by": user}},
            upsert=True
        )
        
        # Mark line as synced
        await db.roster_lines.update_one(
            {"line_id": line["line_id"]},
            {"$set": {"attendance_synced": True, "attendance_sync_at": now, "updated_at": now}}
        )
        synced += 1
    
    await db.roster_audit.insert_one({
        "roster_id": roster_id, "action": "attendance_synced",
        "details": f"Synced {synced} confirmed entries to international attendance",
        "by": user, "at": now
    })
    
    return {"success": True, "synced": synced, "message": f"Synced {synced} entries to attendance"}


# =======================================
# LOCK / UNLOCK
# =======================================

@router.post("/roster/lock")
async def lock_roster(req: dict = Body(...)):
    """Lock or unlock a roster week."""
    session = await check_access(req.get("token"))
    if not is_admin(session):
        raise HTTPException(403, "Admin access required to lock/unlock")
    
    roster_id = req.get("roster_id")
    action = req.get("action", "lock")
    now = datetime.now(timezone.utc).isoformat()
    user = session.get("managerName", "Unknown")
    
    new_status = "locked" if action == "lock" else "draft"
    await db.roster_weeks.update_one(
        {"roster_id": roster_id},
        {"$set": {"status": new_status, "updated_by": user, "updated_at": now}}
    )
    
    await db.roster_audit.insert_one({
        "roster_id": roster_id, "action": action,
        "details": f"Roster {action}ed", "by": user, "at": now
    })
    
    return {"success": True, "message": f"Roster {action}ed"}


# =======================================
# DASHBOARD / STATS
# =======================================

@router.post("/roster/dashboard")
async def roster_dashboard(req: dict = Body(...)):
    """Get roster dashboard stats for a week."""
    session = await check_access(req.get("token"))
    center = (req.get("center") or "").upper()
    week_start = req.get("week_start")
    await check_center_access(session, center)
    
    if not week_start:
        week_start = get_week_info()["week_start"]
    
    week_info = get_week_info(week_start)
    roster = await db.roster_weeks.find_one(
        {"center": center, "week_start": week_info["week_start"]}, {"_id": 0}
    )
    
    if not roster:
        return {"success": True, "exists": False, "week_info": week_info, "stats": {}}
    
    lines = await db.roster_lines.find({"roster_id": roster["roster_id"]}, {"_id": 0}).to_list(500)
    
    # Compute stats
    total = len(lines)
    by_status = {}
    by_role = {}
    by_day = {}
    total_hours = 0
    employees_set = set()
    
    for line in lines:
        s = line.get("status", "draft")
        by_status[s] = by_status.get(s, 0) + 1
        r = line.get("duty_role", "Other")
        by_role[r] = by_role.get(r, 0) + 1
        d = line.get("day", "")
        by_day[d] = by_day.get(d, 0) + 1
        total_hours += line.get("working_hours", 0)
        if line.get("employee_name"):
            employees_set.add(line["employee_name"])
    
    return {
        "success": True, "exists": True,
        "roster": roster, "week_info": week_info,
        "stats": {
            "total_shifts": total,
            "confirmed": by_status.get("confirmed", 0),
            "pending": by_status.get("sent", 0) + by_status.get("draft", 0),
            "denied": by_status.get("denied", 0),
            "replaced": by_status.get("replaced", 0),
            "synced": sum(1 for line in lines if line.get("attendance_synced")),
            "total_hours": round(total_hours, 1),
            "unique_employees": len(employees_set),
            "by_status": by_status,
            "by_role": by_role,
            "by_day": by_day,
        }
    }


# =======================================
# AUDIT LOG
# =======================================

@router.post("/roster/audit")
async def get_audit_log(req: dict = Body(...)):
    """Get audit trail for a roster."""
    session = await check_access(req.get("token"))
    roster_id = req.get("roster_id")
    
    logs = await db.roster_audit.find(
        {"roster_id": roster_id}, {"_id": 0}
    ).sort("at", -1).to_list(200)
    
    return {"success": True, "logs": logs}


# =======================================
# EXPORT
# =======================================

@router.post("/roster/export")
async def export_roster(req: dict = Body(...)):
    """Export roster as Excel."""
    session = await check_access(req.get("token"))
    roster_id = req.get("roster_id")
    
    if not roster_id:
        raise HTTPException(400, "roster_id required")
    
    roster = await db.roster_weeks.find_one({"roster_id": roster_id}, {"_id": 0})
    if not roster:
        raise HTTPException(404, "Roster not found")
    
    lines = await db.roster_lines.find(
        {"roster_id": roster_id}, {"_id": 0}
    ).sort([("date", 1), ("planned_in", 1)]).to_list(500)
    
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from fastapi.responses import Response
    
    wb = Workbook()
    ws = wb.active
    ws.title = f"Roster W{roster.get('week_number', '')}"
    
    # Header
    ws.merge_cells("A1:J1")
    ws["A1"] = f"Weekly Roster — {roster.get('center', '')} — {roster.get('week_label', '')}"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = f"Status: {roster.get('status', '').upper()}"
    
    headers = ["Date", "Day", "Shift", "In", "Out", "Hours", "Employee", "Role", "WhatsApp", "Status"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="8B0000")
        cell.alignment = Alignment(horizontal="center")
    
    status_colors = {
        "confirmed": "C8E6C9", "denied": "FFCDD2", "sent": "BBDEFB",
        "draft": "FFF9C4", "replaced": "FFE0B2"
    }
    
    for row_idx, line in enumerate(lines, 5):
        ws.cell(row=row_idx, column=1, value=line.get("date", ""))
        ws.cell(row=row_idx, column=2, value=line.get("day", ""))
        ws.cell(row=row_idx, column=3, value=line.get("shift_type", ""))
        ws.cell(row=row_idx, column=4, value=line.get("planned_in", ""))
        ws.cell(row=row_idx, column=5, value=line.get("planned_out", ""))
        ws.cell(row=row_idx, column=6, value=line.get("working_hours", 0))
        ws.cell(row=row_idx, column=7, value=line.get("employee_name", ""))
        ws.cell(row=row_idx, column=8, value=line.get("duty_role", ""))
        ws.cell(row=row_idx, column=9, value=line.get("whatsapp", ""))
        status = line.get("status", "draft")
        cell = ws.cell(row=row_idx, column=10, value=status.upper())
        if status in status_colors:
            cell.fill = PatternFill("solid", fgColor=status_colors[status])
    
    # Auto-width (skip merged cells)
    for col_idx in range(1, len(headers) + 1):
        max_len = 0
        col_letter = chr(64 + col_idx) if col_idx <= 26 else "A"
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, min_row=4, values_only=True):
            for val in row:
                max_len = max(max_len, len(str(val or "")))
        ws.column_dimensions[col_letter].width = min(max_len + 3, 25)
    
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    
    filename = f"Roster_{roster.get('center', '')}_{roster.get('week_start', '')}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )
