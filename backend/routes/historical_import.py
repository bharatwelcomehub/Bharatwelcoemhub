"""Historical Data Import

Admin endpoint for ingesting WC-rollup Excel files into a
`historical_monthly_summary` collection. The collection is read by the
WC-Breakdown aggregator and MIS overview as a fallback source of truth
for months that have no daily rows in `daily_sales` / `db.expenses`.
"""
import io
import logging
from datetime import datetime, timezone
from typing import Optional

from fastapi import APIRouter, HTTPException, UploadFile, File, Form
from motor.motor_asyncio import AsyncIOMotorClient
import os
import openpyxl

from .attendance_dashboard import get_session, check_super_admin

logger = logging.getLogger(__name__)

_mongo = AsyncIOMotorClient(os.environ["MONGO_URL"])
db = _mongo[os.environ["DB_NAME"]]

router = APIRouter(prefix="/api/historical", tags=["historical"])


# Sheet name (or uppercased keywords in it) → system center code.
# Matching is: if any key is a substring of the sheet name (upper) → map.
SHEET_TO_CENTER = {
    "S-NAGAR": "PB-SN",
    "SAMBHAJI": "PB-SN",
    "S NAGAR": "PB-SN",
    "DOMBIVLI": "PB-DV",
    "DOMBVLI": "PB-DV",
    "KHARADI": "PB-KN",
    "THANE": "PB-TH",
    "HSR": "PB-HSR",
    "HINJAWADI": "PB-HW",
    "HINJAWADE": "PB-HW",
    "BANER": "PB-MGT",  # TODO: confirm with user — using MGT as fallback
}


def _guess_center_from_sheet(sheet_name: str) -> Optional[str]:
    up = sheet_name.upper()
    for key, code in SHEET_TO_CENTER.items():
        if key in up:
            return code
    return None


# Canonical column aliases. Lowercased header → canonical key.
COLUMN_ALIASES = {
    "month": "month",
    "sale": "sale",
    "sales": "sale",
    "expenses": "expenses",
    "expense": "expenses",
    "p/l": "pnl",
    "pnl": "pnl",
    "profit/loss": "pnl",
    "working capital": "opening_wc",
    "bal. wc.": "closing_wc",
    "bal wc": "closing_wc",
    "balance wc": "closing_wc",
    "bal. wc": "closing_wc",
    "closing wc": "closing_wc",
    "diff.of wc.": "diff_wc",
    "diff of wc": "diff_wc",
    "diff wc": "diff_wc",
    "bank cl.bal.": "bank_balance",
    "bank cl bal": "bank_balance",
    "bank balance": "bank_balance",
    "bank bal": "bank_balance",
    "loan": "loan",
    "rev share": "rev_share",
}


def _parse_wc_sheet(ws) -> tuple[Optional[str], list[dict]]:
    """Return (center_code_guess, list_of_month_rows)."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return None, []
    
    center = _guess_center_from_sheet(ws.title)
    
    # Find header row (containing MONTH/SALE/EXPENSES)
    header_idx = None
    header_map = {}
    for ri, row in enumerate(rows[:15]):
        lowered = [str(c).strip().lower() if c is not None else "" for c in row]
        if any("month" == v for v in lowered) and any("sale" == v or "sales" == v for v in lowered):
            header_idx = ri
            for ci, v in enumerate(lowered):
                if v in COLUMN_ALIASES:
                    header_map[COLUMN_ALIASES[v]] = ci
            break
    
    if header_idx is None or "month" not in header_map:
        return center, []
    
    out = []
    for row in rows[header_idx + 1:]:
        mcell = row[header_map["month"]]
        if mcell is None:
            continue
        # Only accept actual datetime cells (skip stray text / totals)
        if not isinstance(mcell, datetime):
            continue
        month = mcell.strftime("%Y-%m")
        rec = {"month": month}
        for key, ci in header_map.items():
            if key == "month":
                continue
            val = row[ci] if ci < len(row) else None
            if isinstance(val, (int, float)):
                rec[key] = float(val)
            elif val in (None, ""):
                rec[key] = 0.0 if key in ("sale", "expenses", "pnl", "opening_wc", "closing_wc", "diff_wc", "bank_balance") else None
            else:
                try:
                    rec[key] = float(str(val).replace(",", ""))
                except (ValueError, TypeError):
                    rec[key] = None
        out.append(rec)
    return center, out


async def _upsert_rows(center: str, rows: list[dict], source: str, user: str):
    if not center or not rows:
        return 0
    now = datetime.now(timezone.utc).isoformat()
    inserted = 0
    for r in rows:
        doc = {
            "center": center,
            "month": r["month"],
            "sale": r.get("sale", 0) or 0,
            "expenses": r.get("expenses", 0) or 0,
            "pnl": r.get("pnl", 0) or 0,
            "opening_wc": r.get("opening_wc", 0) or 0,
            "closing_wc": r.get("closing_wc", 0) or 0,
            "diff_wc": r.get("diff_wc", 0) or 0,
            "bank_balance": r.get("bank_balance", 0) or 0,
            "source": source,
            "updated_at": now,
            "updated_by": user,
        }
        await db.historical_monthly_summary.update_one(
            {"center": center, "month": r["month"]},
            {"$set": doc, "$setOnInsert": {"created_at": now, "created_by": user}},
            upsert=True,
        )
        inserted += 1
    return inserted


@router.post("/import-wc-file")
async def import_wc_file(file: UploadFile = File(...), token: str = Form(...)):
    """Upload a WC-Assessment Excel file. Each sheet whose name maps to a known
    center will be parsed (MONTH, SALE, EXPENSES, P/L, WORKING CAPITAL, BAL. WC., BANK CL.BAL.)
    and upserted into `historical_monthly_summary`.
    Loans sheets can be imported via /import-loans-file.
    """
    session = await get_session(token)
    if not session or not check_super_admin(session):
        raise HTTPException(403, "Only Super Admin can import historical data")
    
    user = session.get("managerName", "Admin")
    raw = await file.read()
    
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    except Exception as exc:
        raise HTTPException(400, f"Could not read workbook: {exc}")
    
    summary = {
        "file": file.filename,
        "sheets_processed": 0,
        "sheets_skipped": [],
        "centers": {},
        "total_rows": 0,
    }
    for sn in wb.sheetnames:
        if "LOAN" in sn.upper():
            continue  # Loans handled separately
        ws = wb[sn]
        center, rows = _parse_wc_sheet(ws)
        if not center:
            summary["sheets_skipped"].append({"sheet": sn, "reason": "unknown center (no mapping)"})
            continue
        if not rows:
            summary["sheets_skipped"].append({"sheet": sn, "reason": "no parseable month rows"})
            continue
        n = await _upsert_rows(center, rows, source=f"wc_import:{file.filename}:{sn}", user=user)
        summary["sheets_processed"] += 1
        summary["centers"][center] = summary["centers"].get(center, 0) + n
        summary["total_rows"] += n
    
    logger.info(f"Historical WC import by {user}: {summary}")
    return {"success": True, **summary}


@router.post("/summary")
async def get_historical_summary(req: dict):
    """List imported historical data grouped by center."""
    token = req.get("token")
    session = await get_session(token)
    if not session:
        raise HTTPException(401, "Invalid token")
    
    rows = await db.historical_monthly_summary.find({}, {"_id": 0}).to_list(5000)
    by_center: dict = {}
    for r in rows:
        c = r["center"]
        by_center.setdefault(c, {"center": c, "months": 0, "earliest": "", "latest": "", "total_sale": 0, "total_expenses": 0})
        by_center[c]["months"] += 1
        m = r["month"]
        if not by_center[c]["earliest"] or m < by_center[c]["earliest"]:
            by_center[c]["earliest"] = m
        if m > by_center[c]["latest"]:
            by_center[c]["latest"] = m
        by_center[c]["total_sale"] += float(r.get("sale", 0))
        by_center[c]["total_expenses"] += float(r.get("expenses", 0))
    
    return {"centers": list(by_center.values()), "total_rows": len(rows)}


@router.post("/clear")
async def clear_historical(req: dict):
    """Danger: wipe all historical_monthly_summary rows. Super Admin only."""
    token = req.get("token")
    session = await get_session(token)
    if not session or not check_super_admin(session):
        raise HTTPException(403, "Only Super Admin can clear historical data")
    
    center = req.get("center")
    if center:
        res = await db.historical_monthly_summary.delete_many({"center": center})
    else:
        res = await db.historical_monthly_summary.delete_many({})
    return {"success": True, "deleted": res.deleted_count}
