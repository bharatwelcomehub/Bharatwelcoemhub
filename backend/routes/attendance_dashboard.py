# =======================================
# Centralized Attendance Dashboard Module
# Admin Panel - Read-only for Admin, Full access for Super Admin
# =======================================

from fastapi import APIRouter, HTTPException, Response
from pydantic import BaseModel
from typing import List, Optional, Dict, Any
from datetime import datetime, timezone, timedelta
import calendar
import logging
from io import BytesIO

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/attendance-dashboard", tags=["Attendance Dashboard"])

# Get DB reference (will be set from main server)
db = None
verify_token = None
verify_token_async_func = None

def set_db(database):
    global db
    db = database

def set_verify_token(func):
    global verify_token
    verify_token = func

def set_verify_token_async(func):
    global verify_token_async_func
    verify_token_async_func = func

async def get_session(token: str):
    """Get session using async verification (with MongoDB fallback)"""
    if verify_token_async_func:
        session = await verify_token_async_func(token)
        if session:
            return session
    return verify_token(token)

# =======================================
# CONSTANTS
# =======================================

ATTENDANCE_STATUS = {
    "P": {"label": "Present", "color": "green", "weight": 1},
    "A": {"label": "Absent", "color": "red", "weight": 0},
    "HD": {"label": "Half Day", "color": "yellow", "weight": 0.5},
    "WO": {"label": "Week Off", "color": "blue", "weight": 1},
    "L": {"label": "Leave", "color": "orange", "weight": 1},
    "LATE": {"label": "Late", "color": "purple", "weight": 1}
}

# =======================================
# ACCESS CONTROL
# =======================================

def check_admin_access(session) -> bool:
    """Check if user has admin/super admin access for viewing"""
    if not session:
        return False
    if session.get("is_super_admin"):
        return True
    if session.get("is_admin"):
        return True
    # Check for specific role
    roles = session.get("roles", {})
    if roles.get("attendance"):
        return True
    return False

def check_super_admin(session) -> bool:
    """Check if user is Super Admin (for edit access)"""
    if not session:
        return False
    return session.get("is_super_admin", False)

def days_in_month(year: int, month: int) -> int:
    return calendar.monthrange(year, month)[1]

# =======================================
# PYDANTIC MODELS
# =======================================

class DashboardRequest(BaseModel):
    token: str
    date: Optional[str] = None  # YYYY-MM-DD, defaults to today
    month: Optional[str] = None  # YYYY-MM for monthly view
    center: Optional[str] = None  # Optional center filter

class CenterDetailRequest(BaseModel):
    token: str
    center: str
    date: Optional[str] = None
    month: Optional[str] = None

class ExportRequest(BaseModel):
    token: str
    center: Optional[str] = None  # None = all centers
    date: Optional[str] = None
    month: Optional[str] = None
    format: str = "excel"  # excel or pdf

class AttendanceEditRequest(BaseModel):
    token: str
    center: str
    date: str
    employee_name: str
    status: str
    notes: Optional[str] = ""

# =======================================
# DASHBOARD ENDPOINTS
# =======================================

@router.post("/summary")
async def get_dashboard_summary(req: DashboardRequest):
    """Get overall attendance summary for all centers or a specific center"""
    session = await get_session(req.token)
    if not session:
        raise HTTPException(401, "Invalid or expired token")
    
    if not check_admin_access(session):
        raise HTTPException(403, "Access denied. Admin or Super Admin required.")
    
    # Determine date
    target_date = req.date or datetime.now().strftime("%Y-%m-%d")
    filter_center = req.center.upper() if req.center else None
    
    # Get all employees (optionally filtered by center)
    emp_query = {"center": filter_center} if filter_center else {}
    all_employees = await db.employees.find(emp_query, {"_id": 0}).to_list(5000)
    total_employees = len(all_employees)
    
    # Get attendance for the date (optionally filtered by center)
    att_query = {"date": target_date}
    if filter_center:
        att_query["center"] = filter_center
    
    attendance_records = await db.attendance.find(att_query, {"_id": 0}).to_list(10000)
    
    # Build attendance map
    att_map = {}
    for a in attendance_records:
        key = f"{a.get('center')}_{a.get('employeeName')}"
        att_map[key] = a.get("status", "")
    
    # Calculate summary
    present = 0
    absent = 0
    half_day = 0
    week_off = 0
    leave = 0
    late = 0
    not_marked = 0
    
    for emp in all_employees:
        emp_name = emp.get("name", "").upper()
        emp_center = emp.get("center", "")
        key = f"{emp_center}_{emp_name}"
        status = att_map.get(key, "")
        
        if status == "P":
            present += 1
        elif status == "A":
            absent += 1
        elif status == "HD":
            half_day += 1
        elif status == "WO":
            week_off += 1
        elif status == "L":
            leave += 1
        elif status == "LATE":
            late += 1
            present += 1  # Late counts as present
        else:
            not_marked += 1
    
    # Calculate attendance percentage (excluding WO and Leave)
    working_employees = total_employees - week_off - leave
    attendance_pct = round((present + half_day * 0.5) / working_employees * 100, 1) if working_employees > 0 else 0
    
    return {
        "date": target_date,
        "center": filter_center or "all",
        "summary": {
            "total_employees": total_employees,
            "present": present,
            "absent": absent,
            "half_day": half_day,
            "week_off": week_off,
            "leave": leave,
            "late": late,
            "not_marked": not_marked,
            "attendance_percentage": attendance_pct
        },
        "is_super_admin": check_super_admin(session),
        "can_edit": check_super_admin(session)
    }

@router.post("/center-breakdown")
async def get_center_breakdown(req: DashboardRequest):
    """Get attendance breakdown by center"""
    session = await get_session(req.token)
    if not session:
        raise HTTPException(401, "Invalid or expired token")
    
    if not check_admin_access(session):
        raise HTTPException(403, "Access denied. Admin or Super Admin required.")
    
    target_date = req.date or datetime.now().strftime("%Y-%m-%d")
    filter_center = req.center.upper() if req.center else None
    
    # Get centers (optionally filtered)
    center_query = {"$or": [{"code": filter_center}, {"center": filter_center}]} if filter_center else {}
    centers = await db.centers.find(center_query, {"_id": 0}).to_list(100)
    
    # Get all employees grouped by center
    emp_query = {"center": filter_center} if filter_center else {}
    all_employees = await db.employees.find(emp_query, {"_id": 0}).to_list(5000)
    employees_by_center = {}
    for emp in all_employees:
        center = emp.get("center", "")
        if center not in employees_by_center:
            employees_by_center[center] = []
        employees_by_center[center].append(emp)
    
    # Get attendance for the date
    att_query = {"date": target_date}
    if filter_center:
        att_query["center"] = filter_center
    attendance_records = await db.attendance.find(att_query, {"_id": 0}).to_list(10000)
    
    # Build attendance map
    att_map = {}
    for a in attendance_records:
        key = f"{a.get('center')}_{a.get('employeeName')}"
        att_map[key] = a.get("status", "")
    
    # Calculate per-center breakdown
    center_data = []
    for center in centers:
        center_code = center.get("code", center.get("center", ""))
        center_name = center.get("name", center_code)
        center_emps = employees_by_center.get(center_code, [])
        
        total = len(center_emps)
        present = 0
        absent = 0
        half_day = 0
        week_off = 0
        leave = 0
        late = 0
        not_marked = 0
        
        for emp in center_emps:
            emp_name = emp.get("name", "").upper()
            key = f"{center_code}_{emp_name}"
            status = att_map.get(key, "")
            
            if status == "P":
                present += 1
            elif status == "A":
                absent += 1
            elif status == "HD":
                half_day += 1
            elif status == "WO":
                week_off += 1
            elif status == "L":
                leave += 1
            elif status == "LATE":
                late += 1
                present += 1
            else:
                not_marked += 1
        
        working = total - week_off - leave
        att_pct = round((present + half_day * 0.5) / working * 100, 1) if working > 0 else 0
        
        # Determine alert status
        alert = "green"
        if att_pct < 70:
            alert = "red"
        elif att_pct < 85:
            alert = "yellow"
        
        center_data.append({
            "center_code": center_code,
            "center_name": center_name,
            "total_staff": total,
            "present": present,
            "absent": absent,
            "half_day": half_day,
            "week_off": week_off,
            "leave": leave,
            "late": late,
            "not_marked": not_marked,
            "attendance_percentage": att_pct,
            "alert": alert
        })
    
    # Sort by attendance percentage (lowest first for alerts)
    center_data.sort(key=lambda x: x["attendance_percentage"])
    
    return {
        "date": target_date,
        "centers": center_data,
        "is_super_admin": check_super_admin(session)
    }

@router.post("/center-detail")
async def get_center_detail(req: CenterDetailRequest):
    """Get detailed attendance for a specific center"""
    session = await get_session(req.token)
    if not session:
        raise HTTPException(401, "Invalid or expired token")
    
    if not check_admin_access(session):
        raise HTTPException(403, "Access denied. Admin or Super Admin required.")
    
    target_date = req.date or datetime.now().strftime("%Y-%m-%d")
    center = req.center.upper()
    
    # Get employees for this center
    employees = await db.employees.find(
        {"center": center},
        {"_id": 0}
    ).sort("name", 1).to_list(500)
    
    # Get attendance for the date
    attendance_records = await db.attendance.find(
        {"date": target_date, "center": center},
        {"_id": 0}
    ).to_list(1000)
    
    # Build attendance map
    att_map = {}
    for a in attendance_records:
        att_map[a.get("employeeName", "").upper()] = {
            "status": a.get("status", ""),
            "notes": a.get("notes", ""),
            "timestamp": a.get("timestamp", ""),
            "submittedBy": a.get("submittedByMobile", "")
        }
    
    # Build employee list with attendance
    employee_data = []
    for emp in employees:
        emp_name = emp.get("name", "").upper()
        att = att_map.get(emp_name, {})
        
        status = att.get("status", "")
        status_info = ATTENDANCE_STATUS.get(status, {"label": "Not Marked", "color": "gray"})
        
        employee_data.append({
            "name": emp_name,
            "designation": emp.get("designation", ""),
            "status": status,
            "status_label": status_info["label"],
            "status_color": status_info["color"],
            "notes": att.get("notes", ""),
            "timestamp": att.get("timestamp", ""),
            "submitted_by": att.get("submittedBy", "")
        })
    
    # Calculate summary for this center
    total = len(employees)
    present = sum(1 for e in employee_data if e["status"] in ["P", "LATE"])
    absent = sum(1 for e in employee_data if e["status"] == "A")
    half_day = sum(1 for e in employee_data if e["status"] == "HD")
    week_off = sum(1 for e in employee_data if e["status"] == "WO")
    leave = sum(1 for e in employee_data if e["status"] == "L")
    late = sum(1 for e in employee_data if e["status"] == "LATE")
    not_marked = sum(1 for e in employee_data if e["status"] == "")
    
    working = total - week_off - leave
    att_pct = round((present + half_day * 0.5) / working * 100, 1) if working > 0 else 0
    
    return {
        "date": target_date,
        "center": center,
        "summary": {
            "total": total,
            "present": present,
            "absent": absent,
            "half_day": half_day,
            "week_off": week_off,
            "leave": leave,
            "late": late,
            "not_marked": not_marked,
            "attendance_percentage": att_pct
        },
        "employees": employee_data,
        "is_super_admin": check_super_admin(session),
        "can_edit": check_super_admin(session)
    }

# =======================================
# GRID VIEW ENDPOINTS (Excel-like display)
# =======================================

@router.post("/monthly-grid")
async def get_monthly_grid(req: DashboardRequest):
    """Get monthly attendance in Excel-like grid format"""
    session = await get_session(req.token)
    if not session:
        raise HTTPException(401, "Invalid or expired token")
    
    if not check_admin_access(session):
        raise HTTPException(403, "Access denied. Admin or Super Admin required.")
    
    # Determine month
    month = req.month or datetime.now().strftime("%Y-%m")
    filter_center = req.center.upper() if req.center else None
    
    year, mon = map(int, month.split("-"))
    dim = days_in_month(year, mon)
    
    # Get employees (optionally filtered by center)
    emp_query = {"center": filter_center} if filter_center else {}
    employees = await db.employees.find(emp_query, {"_id": 0}).sort([("center", 1), ("name", 1)]).to_list(5000)
    
    # Get attendance for the month
    start_date = f"{month}-01"
    end_date = f"{month}-{dim:02d}"
    
    att_query = {"date": {"$gte": start_date, "$lte": end_date}}
    if filter_center:
        att_query["center"] = filter_center
    
    attendance_records = await db.attendance.find(att_query, {"_id": 0}).to_list(50000)
    
    # Build attendance map: {center_employee_date: status}
    att_map = {}
    for a in attendance_records:
        key = f"{a.get('center')}_{a.get('employeeName')}_{a.get('date')}"
        att_map[key] = a.get("status", "")
    
    # Build employee grid data
    employee_grid = []
    center_stats = {}
    
    for emp in employees:
        emp_name = emp.get("name", "").upper()
        emp_center = emp.get("center", "")
        
        # Initialize center stats
        if emp_center not in center_stats:
            center_stats[emp_center] = {
                "total_staff": 0, "present": 0, "absent": 0, 
                "half_day": 0, "week_off": 0, "leave": 0
            }
        center_stats[emp_center]["total_staff"] += 1
        
        # Get attendance for each day
        attendance = []
        for d in range(1, dim + 1):
            date_str = f"{month}-{d:02d}"
            key = f"{emp_center}_{emp_name}_{date_str}"
            status = att_map.get(key, "")
            attendance.append(status)
            
            # Update center stats
            if status == "P" or status == "LATE":
                center_stats[emp_center]["present"] += 1
            elif status == "A":
                center_stats[emp_center]["absent"] += 1
            elif status == "HD":
                center_stats[emp_center]["half_day"] += 1
            elif status == "WO":
                center_stats[emp_center]["week_off"] += 1
            elif status == "L":
                center_stats[emp_center]["leave"] += 1
        
        employee_grid.append({
            "name": emp_name,
            "center": emp_center,
            "designation": emp.get("designation", ""),
            "attendance": attendance
        })
    
    # Build center summary
    center_summary = []
    for center_code, stats in center_stats.items():
        working_days = stats["total_staff"] * dim - stats["week_off"] - stats["leave"]
        att_pct = round((stats["present"] + stats["half_day"] * 0.5) / working_days * 100, 1) if working_days > 0 else 0
        center_summary.append({
            "center": center_code,
            **stats,
            "attendance_pct": att_pct
        })
    
    center_summary.sort(key=lambda x: x["attendance_pct"])
    
    return {
        "month": month,
        "days_in_month": dim,
        "employees": employee_grid,
        "center_summary": center_summary,
        "is_super_admin": check_super_admin(session)
    }

@router.post("/daily-grid")
async def get_daily_grid(req: DashboardRequest):
    """Get daily attendance in grid format"""
    session = await get_session(req.token)
    if not session:
        raise HTTPException(401, "Invalid or expired token")
    
    if not check_admin_access(session):
        raise HTTPException(403, "Access denied. Admin or Super Admin required.")
    
    target_date = req.date or datetime.now().strftime("%Y-%m-%d")
    filter_center = req.center.upper() if req.center else None
    
    # Get employees (optionally filtered by center)
    emp_query = {"center": filter_center} if filter_center else {}
    employees = await db.employees.find(emp_query, {"_id": 0}).sort([("center", 1), ("name", 1)]).to_list(5000)
    
    # Get attendance for the date
    att_query = {"date": target_date}
    if filter_center:
        att_query["center"] = filter_center
    
    attendance_records = await db.attendance.find(att_query, {"_id": 0}).to_list(10000)
    
    # Build attendance map
    att_map = {}
    for a in attendance_records:
        key = f"{a.get('center')}_{a.get('employeeName')}"
        att_map[key] = {
            "status": a.get("status", ""),
            "notes": a.get("notes", "")
        }
    
    # Build employee data
    employee_data = []
    for emp in employees:
        emp_name = emp.get("name", "").upper()
        emp_center = emp.get("center", "")
        key = f"{emp_center}_{emp_name}"
        att = att_map.get(key, {"status": "", "notes": ""})
        
        employee_data.append({
            "name": emp_name,
            "center": emp_center,
            "designation": emp.get("designation", ""),
            "status": att["status"],
            "notes": att["notes"]
        })
    
    return {
        "date": target_date,
        "employees": employee_data,
        "is_super_admin": check_super_admin(session)
    }

@router.post("/monthly-trend")
async def get_monthly_trend(req: DashboardRequest):
    """Get attendance trend for a month"""
    session = await get_session(req.token)
    if not session:
        raise HTTPException(401, "Invalid or expired token")
    
    if not check_admin_access(session):
        raise HTTPException(403, "Access denied. Admin or Super Admin required.")
    
    # Determine month
    if req.month:
        month = req.month
    else:
        month = datetime.now().strftime("%Y-%m")
    
    year, mon = map(int, month.split("-"))
    dim = days_in_month(year, mon)
    
    # Get all employees
    all_employees = await db.employees.find({}, {"_id": 0}).to_list(5000)
    total_employees = len(all_employees)
    
    # Get attendance for the month
    start_date = f"{month}-01"
    end_date = f"{month}-{dim:02d}"
    
    attendance_records = await db.attendance.find(
        {"date": {"$gte": start_date, "$lte": end_date}},
        {"_id": 0}
    ).to_list(50000)
    
    # Group by date
    daily_att = {}
    for a in attendance_records:
        date = a.get("date", "")
        if date not in daily_att:
            daily_att[date] = {"P": 0, "A": 0, "HD": 0, "WO": 0, "L": 0, "LATE": 0}
        status = a.get("status", "")
        if status in daily_att[date]:
            daily_att[date][status] += 1
    
    # Build daily trend
    trend = []
    for d in range(1, dim + 1):
        date_str = f"{month}-{d:02d}"
        day_data = daily_att.get(date_str, {"P": 0, "A": 0, "HD": 0, "WO": 0, "L": 0, "LATE": 0})
        
        present = day_data["P"] + day_data["LATE"]
        working = total_employees - day_data["WO"] - day_data["L"]
        att_pct = round((present + day_data["HD"] * 0.5) / working * 100, 1) if working > 0 else 0
        
        trend.append({
            "date": date_str,
            "day": d,
            "present": present,
            "absent": day_data["A"],
            "half_day": day_data["HD"],
            "week_off": day_data["WO"],
            "leave": day_data["L"],
            "late": day_data["LATE"],
            "attendance_percentage": att_pct
        })
    
    return {
        "month": month,
        "total_employees": total_employees,
        "days_in_month": dim,
        "trend": trend
    }

@router.post("/center-comparison")
async def get_center_comparison(req: DashboardRequest):
    """Get attendance comparison across centers for charts"""
    session = await get_session(req.token)
    if not session:
        raise HTTPException(401, "Invalid or expired token")
    
    if not check_admin_access(session):
        raise HTTPException(403, "Access denied. Admin or Super Admin required.")
    
    # Determine month
    if req.month:
        month = req.month
    else:
        month = datetime.now().strftime("%Y-%m")
    
    year, mon = map(int, month.split("-"))
    dim = days_in_month(year, mon)
    
    # Get all centers
    centers = await db.centers.find({}, {"_id": 0}).to_list(100)
    
    # Get all employees grouped by center
    all_employees = await db.employees.find({}, {"_id": 0}).to_list(5000)
    employees_by_center = {}
    for emp in all_employees:
        center = emp.get("center", "")
        employees_by_center[center] = employees_by_center.get(center, 0) + 1
    
    # Get attendance for the month
    start_date = f"{month}-01"
    end_date = f"{month}-{dim:02d}"
    
    attendance_records = await db.attendance.find(
        {"date": {"$gte": start_date, "$lte": end_date}},
        {"_id": 0}
    ).to_list(50000)
    
    # Group by center
    center_att = {}
    for a in attendance_records:
        center = a.get("center", "")
        if center not in center_att:
            center_att[center] = {"P": 0, "A": 0, "HD": 0, "WO": 0, "L": 0, "LATE": 0, "total_days": 0}
        status = a.get("status", "")
        if status in center_att[center]:
            center_att[center][status] += 1
        center_att[center]["total_days"] += 1
    
    # Calculate per-center stats
    comparison = []
    for center in centers:
        center_code = center.get("code", center.get("center", ""))
        center_name = center.get("name", center_code)
        
        total_emps = employees_by_center.get(center_code, 0)
        data = center_att.get(center_code, {"P": 0, "A": 0, "HD": 0, "WO": 0, "L": 0, "LATE": 0})
        
        present = data["P"] + data["LATE"]
        total_working_days = (total_emps * dim) - data["WO"] - data["L"]
        att_pct = round((present + data["HD"] * 0.5) / total_working_days * 100, 1) if total_working_days > 0 else 0
        
        comparison.append({
            "center_code": center_code,
            "center_name": center_name,
            "total_employees": total_emps,
            "total_present": present,
            "total_absent": data["A"],
            "attendance_percentage": att_pct
        })
    
    # Sort by attendance
    comparison.sort(key=lambda x: x["attendance_percentage"], reverse=True)
    
    return {
        "month": month,
        "comparison": comparison
    }

# =======================================
# EXPORT ENDPOINTS
# =======================================

@router.post("/export")
async def export_attendance(req: ExportRequest):
    """Export attendance report"""
    session = await get_session(req.token)
    if not session:
        raise HTTPException(401, "Invalid or expired token")
    
    if not check_admin_access(session):
        raise HTTPException(403, "Access denied. Admin or Super Admin required.")
    
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        
        target_date = req.date or datetime.now().strftime("%Y-%m-%d")
        
        # Get employees
        query = {}
        if req.center:
            query["center"] = req.center.upper()
        
        employees = await db.employees.find(query, {"_id": 0}).sort([("center", 1), ("name", 1)]).to_list(5000)
        
        # Get attendance
        att_query = {"date": target_date}
        if req.center:
            att_query["center"] = req.center.upper()
        
        attendance_records = await db.attendance.find(att_query, {"_id": 0}).to_list(10000)
        
        # Build attendance map
        att_map = {}
        for a in attendance_records:
            key = f"{a.get('center')}_{a.get('employeeName')}"
            att_map[key] = {
                "status": a.get("status", ""),
                "notes": a.get("notes", "")
            }
        
        # Create workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance Report"
        
        # Styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="8B0000", end_color="8B0000", fill_type="solid")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Title
        ws.merge_cells('A1:F1')
        ws['A1'] = f"Purnabramha Attendance Report - {target_date}"
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = Alignment(horizontal='center')
        
        # Headers
        headers = ["Center", "Employee Name", "Designation", "Status", "Status Label", "Notes"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(horizontal='center')
        
        # Data
        row = 4
        for emp in employees:
            emp_name = emp.get("name", "").upper()
            center = emp.get("center", "")
            key = f"{center}_{emp_name}"
            att = att_map.get(key, {"status": "", "notes": ""})
            
            status = att["status"]
            status_info = ATTENDANCE_STATUS.get(status, {"label": "Not Marked"})
            
            ws.cell(row=row, column=1, value=center).border = border
            ws.cell(row=row, column=2, value=emp_name).border = border
            ws.cell(row=row, column=3, value=emp.get("designation", "")).border = border
            ws.cell(row=row, column=4, value=status).border = border
            ws.cell(row=row, column=5, value=status_info["label"]).border = border
            ws.cell(row=row, column=6, value=att["notes"]).border = border
            
            row += 1
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 12
        ws.column_dimensions['B'].width = 25
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 10
        ws.column_dimensions['E'].width = 15
        ws.column_dimensions['F'].width = 30
        
        # Save to bytes
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        filename = f"Attendance_{req.center or 'ALL'}_{target_date}.xlsx"
        
        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )
        
    except Exception as e:
        logger.error(f"Export error: {e}")
        raise HTTPException(500, str(e))

@router.post("/export-monthly")
async def export_monthly_attendance(req: ExportRequest):
    """Export monthly attendance report"""
    session = await get_session(req.token)
    if not session:
        raise HTTPException(401, "Invalid or expired token")
    
    if not check_admin_access(session):
        raise HTTPException(403, "Access denied. Admin or Super Admin required.")
    
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        
        # Determine month
        if req.month:
            month = req.month
        else:
            month = datetime.now().strftime("%Y-%m")
        
        year, mon = map(int, month.split("-"))
        dim = days_in_month(year, mon)
        
        # Get employees
        query = {}
        if req.center:
            query["center"] = req.center.upper()
        
        employees = await db.employees.find(query, {"_id": 0}).sort([("center", 1), ("name", 1)]).to_list(5000)
        
        # Get attendance for the month
        start_date = f"{month}-01"
        end_date = f"{month}-{dim:02d}"
        
        att_query = {"date": {"$gte": start_date, "$lte": end_date}}
        if req.center:
            att_query["center"] = req.center.upper()
        
        attendance_records = await db.attendance.find(att_query, {"_id": 0}).to_list(50000)
        
        # Build attendance map
        att_map = {}
        for a in attendance_records:
            key = f"{a.get('center')}_{a.get('employeeName')}_{a.get('date')}"
            att_map[key] = a.get("status", "")
        
        # Create workbook
        wb = Workbook()
        ws = wb.active
        ws.title = f"Attendance {month}"
        
        # Styles
        header_font = Font(bold=True, color="FFFFFF", size=9)
        header_fill = PatternFill(start_color="8B0000", end_color="8B0000", fill_type="solid")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        status_fills = {
            "P": PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid"),
            "A": PatternFill(start_color="FFB6C1", end_color="FFB6C1", fill_type="solid"),
            "HD": PatternFill(start_color="FFFFE0", end_color="FFFFE0", fill_type="solid"),
            "WO": PatternFill(start_color="ADD8E6", end_color="ADD8E6", fill_type="solid"),
            "L": PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid"),
            "LATE": PatternFill(start_color="DDA0DD", end_color="DDA0DD", fill_type="solid")
        }
        
        # Title
        ws.merge_cells(f'A1:AH1')
        ws['A1'] = f"Purnabramha Monthly Attendance - {month}"
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = Alignment(horizontal='center')
        
        # Headers
        headers = ["Center", "Employee", "Designation"] + [str(d) for d in range(1, dim + 1)] + ["Present", "Absent", "HD", "WO", "Leave", "Late", "%"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(horizontal='center')
        
        # Data
        row = 4
        for emp in employees:
            emp_name = emp.get("name", "").upper()
            center = emp.get("center", "")
            
            ws.cell(row=row, column=1, value=center).border = border
            ws.cell(row=row, column=2, value=emp_name).border = border
            ws.cell(row=row, column=3, value=emp.get("designation", "")).border = border
            
            # Daily status
            counts = {"P": 0, "A": 0, "HD": 0, "WO": 0, "L": 0, "LATE": 0}
            for d in range(1, dim + 1):
                date_str = f"{month}-{d:02d}"
                key = f"{center}_{emp_name}_{date_str}"
                status = att_map.get(key, "")
                
                cell = ws.cell(row=row, column=3 + d, value=status)
                cell.border = border
                cell.alignment = Alignment(horizontal='center')
                
                if status in status_fills:
                    cell.fill = status_fills[status]
                    counts[status] += 1
            
            # Summary columns
            col_offset = 3 + dim
            ws.cell(row=row, column=col_offset + 1, value=counts["P"] + counts["LATE"]).border = border
            ws.cell(row=row, column=col_offset + 2, value=counts["A"]).border = border
            ws.cell(row=row, column=col_offset + 3, value=counts["HD"]).border = border
            ws.cell(row=row, column=col_offset + 4, value=counts["WO"]).border = border
            ws.cell(row=row, column=col_offset + 5, value=counts["L"]).border = border
            ws.cell(row=row, column=col_offset + 6, value=counts["LATE"]).border = border
            
            working_days = dim - counts["WO"] - counts["L"]
            att_pct = round((counts["P"] + counts["LATE"] + counts["HD"] * 0.5) / working_days * 100, 1) if working_days > 0 else 0
            ws.cell(row=row, column=col_offset + 7, value=f"{att_pct}%").border = border
            
            row += 1
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 12
        for d in range(1, dim + 1):
            ws.column_dimensions[chr(67 + d) if d <= 23 else 'A' + chr(67 + d - 26)].width = 4
        
        # Save to bytes
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        filename = f"Monthly_Attendance_{req.center or 'ALL'}_{month}.xlsx"
        
        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )
        
    except Exception as e:
        logger.error(f"Monthly export error: {e}")
        raise HTTPException(500, str(e))

# =======================================
# SUPER ADMIN EDIT ENDPOINTS
# =======================================

@router.post("/edit")
async def edit_attendance(req: AttendanceEditRequest):
    """Edit attendance (Super Admin only)"""
    session = await get_session(req.token)
    if not session:
        raise HTTPException(401, "Invalid or expired token")
    
    if not check_super_admin(session):
        raise HTTPException(403, "Only Super Admin can edit attendance")
    
    # Validate status
    if req.status not in ATTENDANCE_STATUS and req.status != "":
        raise HTTPException(400, f"Invalid status. Must be one of: {', '.join(ATTENDANCE_STATUS.keys())}")
    
    # Update attendance
    doc = {
        "date": req.date,
        "center": req.center.upper(),
        "employeeName": req.employee_name.upper(),
        "status": req.status.upper(),
        "notes": req.notes or "",
        "submittedByMobile": session.get("mobile", ""),
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "editedBy": session.get("managerName", "Super Admin"),
        "editedAt": datetime.now(timezone.utc).isoformat()
    }
    
    result = await db.attendance.update_one(
        {"date": req.date, "center": req.center.upper(), "employeeName": req.employee_name.upper()},
        {"$set": doc},
        upsert=True
    )
    
    # Log the edit
    await db.attendance_audit.insert_one({
        "action": "EDIT",
        "date": req.date,
        "center": req.center.upper(),
        "employee": req.employee_name.upper(),
        "new_status": req.status,
        "edited_by": session.get("managerName", ""),
        "edited_at": datetime.now(timezone.utc).isoformat()
    })
    
    logger.info(f"Attendance edited: {req.center}/{req.employee_name} on {req.date} to {req.status} by {session.get('managerName')}")
    
    return {"success": True, "message": "Attendance updated"}

@router.get("/status-options")
async def get_status_options():
    """Get available attendance status options"""
    return {
        "statuses": [
            {"code": code, **info}
            for code, info in ATTENDANCE_STATUS.items()
        ]
    }
